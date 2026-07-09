import re
import html
from datetime import datetime

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
import docx
from docx.shared import RGBColor, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ---------------------------------------------------------------------------
# Markdown → ReportLab helpers
# ---------------------------------------------------------------------------

def _rl_escape(text: str) -> str:
    """Escape XML special characters for ReportLab, then apply inline markdown."""
    escaped = html.escape(text)
    # Bold **text** and __text__
    escaped = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', escaped)
    escaped = re.sub(r'__(.*?)__', r'<b>\1</b>', escaped)
    # Italic *text* and _text_
    escaped = re.sub(r'\*(.*?)\*', r'<i>\1</i>', escaped)
    escaped = re.sub(r'_((?!_).*?)_', r'<i>\1</i>', escaped)
    # Inline code `text`
    escaped = re.sub(r'`(.*?)`', r'<font face="Courier" backColor="#f1f5f9">\1</font>', escaped)
    return escaped


def _parse_md_table(lines: list[str]) -> list[list[str]] | None:
    """Try to parse a markdown table block into a list of rows (each a list of cell strings)."""
    rows = []
    for line in lines:
        if re.match(r'\|?\s*[-:]+[-| :]+\s*\|?', line.strip()):
            # separator row — skip
            continue
        cells = [c.strip() for c in re.split(r'\|', line.strip()) if c.strip() != '']
        if cells:
            rows.append(cells)
    return rows if len(rows) >= 1 else None


def _flowables_from_markdown(content: str, styles: dict) -> list:
    """
    Convert markdown content to a list of ReportLab Flowables.
    Handles: headings, bullet lists, numbered lists, tables, code blocks, and paragraphs.
    """
    flowables = []
    lines = content.split('\n')
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # --- Code block ---
        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            code_text = '\n'.join(code_lines)
            # Render as a shaded paragraph
            flowables.append(Spacer(1, 4))
            for cl in code_text.split('\n'):
                escaped = html.escape(cl) if cl.strip() else '&nbsp;'
                flowables.append(Paragraph(escaped, styles['code']))
            flowables.append(Spacer(1, 4))
            i += 1
            continue

        # --- Markdown table ---
        if '|' in stripped and stripped.startswith('|'):
            table_lines = []
            while i < len(lines) and '|' in lines[i]:
                table_lines.append(lines[i])
                i += 1
            rows = _parse_md_table(table_lines)
            if rows:
                flowables.append(_build_rl_table(rows, styles))
                flowables.append(Spacer(1, 8))
            continue

        # --- Heading h1-h4 ---
        heading_match = re.match(r'^(#{1,4})\s+(.*)', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = _rl_escape(heading_match.group(2))
            style_name = f'h{min(level, 4)}'
            flowables.append(Spacer(1, 6))
            flowables.append(Paragraph(text, styles[style_name]))
            flowables.append(Spacer(1, 2))
            i += 1
            continue

        # --- Horizontal rule ---
        if re.match(r'^[-*_]{3,}$', stripped):
            flowables.append(Spacer(1, 4))
            flowables.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#cbd5e1')))
            flowables.append(Spacer(1, 4))
            i += 1
            continue

        # --- Bullet list item ---
        bullet_match = re.match(r'^[-*+]\s+(.*)', stripped)
        if bullet_match:
            text = _rl_escape(bullet_match.group(1))
            flowables.append(Paragraph(f'• {text}', styles['bullet']))
            i += 1
            continue

        # --- Numbered list item ---
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            text = _rl_escape(num_match.group(2))
            flowables.append(Paragraph(f'{num_match.group(1)}. {text}', styles['bullet']))
            i += 1
            continue

        # --- Blank line ---
        if not stripped:
            flowables.append(Spacer(1, 5))
            i += 1
            continue

        # --- Regular paragraph ---
        text = _rl_escape(stripped)
        flowables.append(Paragraph(text, styles['body']))
        i += 1

    return flowables


def _build_rl_table(rows: list[list[str]], styles: dict):
    """Build a ReportLab Table from a list of rows."""
    num_cols = max(len(r) for r in rows)
    # Pad rows with empty cells if needed
    padded = [r + [''] * (num_cols - len(r)) for r in rows]

    table_data = []
    for ri, row in enumerate(padded):
        table_row = []
        for ci, cell in enumerate(row):
            style = styles['table_header'] if ri == 0 else styles['table_cell']
            table_row.append(Paragraph(_rl_escape(cell), style))
        table_data.append(table_row)

    col_width = (letter[0] - 108) / num_cols  # 108 = L+R margins
    t = Table(table_data, colWidths=[col_width] * num_cols, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f0f9ff'), colors.white]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    return t


# ---------------------------------------------------------------------------
# PDF Export
# ---------------------------------------------------------------------------

def generate_pdf_export(messages: list[dict], title: str, assoc_file: str | None, output_path: str):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54
    )

    base = getSampleStyleSheet()

    styles = {
        'title': ParagraphStyle('DocTitle', parent=base['Heading1'],
            fontSize=20, leading=24, textColor=colors.HexColor('#1a202c'), spaceAfter=6),
        'meta': ParagraphStyle('DocMeta', parent=base['Normal'],
            fontSize=9, leading=12, textColor=colors.HexColor('#718096'), spaceAfter=12),
        'role_user': ParagraphStyle('RoleUser', parent=base['Normal'],
            fontSize=9, fontName='Helvetica-Bold',
            textColor=colors.HexColor('#2563eb'), spaceAfter=3),
        'role_assistant': ParagraphStyle('RoleAssist', parent=base['Normal'],
            fontSize=9, fontName='Helvetica-Bold',
            textColor=colors.HexColor('#0f766e'), spaceAfter=3),
        'body': ParagraphStyle('Body', parent=base['Normal'],
            fontSize=10, leading=14, textColor=colors.HexColor('#1e293b'), spaceAfter=3, leftIndent=12),
        'bullet': ParagraphStyle('Bullet', parent=base['Normal'],
            fontSize=10, leading=13, textColor=colors.HexColor('#1e293b'),
            leftIndent=24, spaceAfter=3),
        'code': ParagraphStyle('Code', parent=base['Normal'],
            fontName='Courier', fontSize=8.5, leading=12,
            backColor=colors.HexColor('#f1f5f9'), leftIndent=16,
            textColor=colors.HexColor('#1e293b'), spaceAfter=1),
        'h1': ParagraphStyle('H1', parent=base['Heading2'],
            fontSize=14, textColor=colors.HexColor('#1e3a8a'), leftIndent=12, spaceBefore=4),
        'h2': ParagraphStyle('H2', parent=base['Heading3'],
            fontSize=12, textColor=colors.HexColor('#1e3a8a'), leftIndent=12, spaceBefore=4),
        'h3': ParagraphStyle('H3', parent=base['Heading4'],
            fontSize=11, textColor=colors.HexColor('#334155'), leftIndent=12, spaceBefore=3),
        'h4': ParagraphStyle('H4', parent=base['Normal'],
            fontSize=10, fontName='Helvetica-Bold',
            textColor=colors.HexColor('#334155'), leftIndent=12, spaceBefore=2),
        'table_header': ParagraphStyle('TblHdr', parent=base['Normal'],
            fontSize=9, fontName='Helvetica-Bold', textColor=colors.white),
        'table_cell': ParagraphStyle('TblCell', parent=base['Normal'],
            fontSize=9, textColor=colors.HexColor('#1e293b')),
        'user_body': ParagraphStyle('UserBody', parent=base['Normal'],
            fontSize=10, leading=14, textColor=colors.HexColor('#1e3a8a'), leftIndent=12),
    }

    story = []
    story.append(Paragraph(title, styles['title']))
    meta_text = f"Exported: {datetime.now().strftime('%B %d, %Y %H:%M')}"
    if assoc_file:
        meta_text += f"  |  Document: {assoc_file}"
    story.append(Paragraph(meta_text, styles['meta']))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e2e8f0')))
    story.append(Spacer(1, 12))

    for msg in messages:
        role = msg.get("role", "user")
        content = msg.get("content", "")

        if role == "user":
            role_p = Paragraph("YOU", styles['role_user'])
            # User messages are usually short — table is fine
            content_p = Paragraph(_rl_escape(content), styles['user_body'])
            t = Table([[role_p], [content_p]], colWidths=[doc.width])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f0f9ff')),
                ('TOPPADDING', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                ('LEFTPADDING', (0, 0), (-1, -1), 12),
                ('RIGHTPADDING', (0, 0), (-1, -1), 12),
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#bae6fd')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))
            story.append(KeepTogether(t))
        else:
            # Assistant: render role label then flowing, fully-formatted content
            story.append(Paragraph("DEEPINSIGHT", styles['role_assistant']))
            story.extend(_flowables_from_markdown(content, styles))

        story.append(Spacer(1, 14))

    doc.build(story)


# ---------------------------------------------------------------------------
# DOCX Export
# ---------------------------------------------------------------------------

def _apply_inline_md_to_run_group(paragraph, text: str):
    """Parse a single line of inline markdown and add styled runs to a docx paragraph."""
    # Pattern: **bold**, *italic*, `code`, plain text
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    parts = pattern.split(text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def _add_md_table_to_doc(doc, table_lines: list[str]):
    rows = _parse_md_table(table_lines)
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'
    for ri, row in enumerate(rows):
        for ci in range(num_cols):
            cell_text = row[ci] if ci < len(row) else ''
            cell = table.cell(ri, ci)
            cell.text = cell_text
            run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else None
            if ri == 0:
                cell.paragraphs[0].runs[0].bold = True if run else None
                # Header shading
                from docx.oxml.ns import qn
                from docx.oxml import OxmlElement
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), '1e3a8a')
                tcPr.append(shd)
                if cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _render_md_to_docx(doc, content: str):
    """Convert markdown content to DOCX elements."""
    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code block
        if stripped.startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Courier New'
            run.font.size = Pt(8.5)
            p.paragraph_format.left_indent = Inches(0.25)
            i += 1
            continue

        # Markdown table
        if '|' in stripped and stripped.startswith('|'):
            table_lines = []
            while i < len(lines) and '|' in lines[i]:
                table_lines.append(lines[i])
                i += 1
            _add_md_table_to_doc(doc, table_lines)
            continue

        # Headings
        heading_match = re.match(r'^(#{1,4})\s+(.*)', stripped)
        if heading_match:
            level = min(len(heading_match.group(1)) + 1, 4)
            doc.add_heading(heading_match.group(2), level=level)
            i += 1
            continue

        # HR
        if re.match(r'^[-*_]{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.border_bottom = True
            i += 1
            continue

        # Bullet
        bullet_match = re.match(r'^[-*+]\s+(.*)', stripped)
        if bullet_match:
            p = doc.add_paragraph(style='List Bullet')
            _apply_inline_md_to_run_group(p, bullet_match.group(1))
            i += 1
            continue

        # Numbered list
        num_match = re.match(r'^\d+\.\s+(.*)', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            _apply_inline_md_to_run_group(p, num_match.group(1))
            i += 1
            continue

        # Blank line
        if not stripped:
            doc.add_paragraph()
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        _apply_inline_md_to_run_group(p, stripped)
        i += 1


def generate_docx_export(messages: list[dict], title: str, assoc_file: str | None, output_path: str):
    doc = docx.Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Title
    title_p = doc.add_heading(title, 0)

    # Metadata
    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(f"Exported: {datetime.now().strftime('%B %d, %Y %H:%M')}")
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    if assoc_file:
        meta_p.add_run(f"   |   Document: {assoc_file}").font.size = Pt(9)

    doc.add_paragraph()

    for msg in messages:
        role = msg.get("role", "user")
        content = msg.get("content", "")

        if role == "user":
            # Role label
            role_p = doc.add_paragraph()
            role_run = role_p.add_run("YOU")
            role_run.bold = True
            role_run.font.size = Pt(9)
            role_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

            # User message (usually short, no complex markdown)
            body_p = doc.add_paragraph()
            body_p.paragraph_format.left_indent = Inches(0.2)
            body_p.add_run(content)
        else:
            # Role label
            role_p = doc.add_paragraph()
            role_run = role_p.add_run("DEEPINSIGHT")
            role_run.bold = True
            role_run.font.size = Pt(9)
            role_run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

            # Render full markdown
            _render_md_to_docx(doc, content)

        doc.add_paragraph()

    doc.save(output_path)
