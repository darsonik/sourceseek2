import base64
import mimetypes
import os

import pdfplumber
import docx
import openpyxl
from openai import OpenAI
from openpyxl.utils import get_column_letter
from PIL import Image
from pydantic import BaseModel, Field
from typing import List, Any

from app.core.config import settings
from app.services.embeddings import get_embeddings


class ParsedChunkPDF(BaseModel):
    """
    Represents a single extracted line of text from a PDF, along with
    its exact location metadata that will be stored in the database.
    """
    content: str
    embedding: List[float]
    location_metadata: dict = Field(
        default_factory=dict,
        description="Exact location of this chunk, e.g. {'page': 3, 'line': 12, 'bbox': {...}}"
    )

class ParsedChunkDOC(BaseModel):
    """
    Represents a single extracted line of text from a DOCX file, along with
    its exact location metadata that will be stored in the database.
    """
    content: str
    embedding: List[float]
    location_metadata: dict = Field(
        default_factory=dict,
        description="Exact location of this chunk, e.g. {'paragraph_index': 4, 'style': 'Heading 1'}"
    )


class ParsedChunkXLSX(BaseModel):
    """
    Represents a single extracted row of data from an Excel file, along with
    its exact location metadata (sheet name, row number, cell references).
    Each row becomes one chunk so there is enough text context for semantic
    embeddings, while the metadata pinpoints each individual cell value.
    """
    content: str
    embedding: List[float]
    location_metadata: dict = Field(
        default_factory=dict,
        description="e.g. {'sheet': 'Sheet1', 'row': 4, 'cells': [{'ref': 'A4', 'column': 'Name', 'value': 'ID1234'}]}"
    )

class ParsedTextFromImages(BaseModel):
    """
    Texts after doing OCR on the images using a Vision AI model.
    """
    content: str
    embedding: List[float]
    location_metadata: dict = Field(
        default_factory=dict,
        description="Which picture contains this text, e.g. {'image_name': 'image1.png'}"
    )


def parse_pdf(file_path: str) -> List[ParsedChunkPDF]:
    """
    Parses a PDF file and returns a list of ParsedChunkPDF objects.
    """
    temp_chunks: List[dict] = []

    with pdfplumber.open(file_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(
                x_tolerance=3,
                y_tolerance=3,
                keep_blank_chars=False,
                use_text_flow=True,  # Respects natural reading order
            )

            if not words:
                continue

            lines: dict[float, list] = {}
            for word in words:
                line_key = round(word["top"] / 2) * 2
                if line_key not in lines:
                    lines[line_key] = []
                lines[line_key].append(word)

            sorted_line_keys = sorted(lines.keys())

            for line_number, line_key in enumerate(sorted_line_keys, start=1):
                line_words = lines[line_key]
                line_words.sort(key=lambda w: w["x0"])
                line_text = " ".join(w["text"] for w in line_words).strip()

                if not line_text:
                    continue

                bbox = {
                    "x0": min(w["x0"] for w in line_words),
                    "top": min(w["top"] for w in line_words),
                    "x1": max(w["x1"] for w in line_words),
                    "bottom": max(w["bottom"] for w in line_words),
                }

                temp_chunks.append({
                    "content": line_text,
                    "location_metadata": {
                        "page": page_number,
                        "line": line_number,
                        "bbox": bbox,
                    }
                })

    if not temp_chunks:
        return []

    # Batch process all collected texts
    texts = [c["content"] for c in temp_chunks]
    embeddings = get_embeddings(texts)

    return [
        ParsedChunkPDF(
            content=c["content"],
            embedding=emb,
            location_metadata=c["location_metadata"]
        )
        for c, emb in zip(temp_chunks, embeddings)
    ]


def parse_docx(file_path: str) -> List[ParsedChunkDOC]:
    """
    Parses a DOCX file and returns a list of ParsedChunkDOC objects.
    """
    temp_chunks: List[dict] = []
    doc = docx.Document(file_path)
    current_section_heading: str | None = None

    for paragraph_index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue

        style_name: str = paragraph.style.name

        if style_name.startswith("Heading"):
            current_section_heading = text

        temp_chunks.append({
            "content": text,
            "location_metadata": {
                "paragraph_index": paragraph_index,
                "style": style_name,
                "section_heading": current_section_heading,
            }
        })

    if not temp_chunks:
        return []

    # Batch process all collected texts
    texts = [c["content"] for c in temp_chunks]
    embeddings = get_embeddings(texts)

    return [
        ParsedChunkDOC(
            content=c["content"],
            embedding=emb,
            location_metadata=c["location_metadata"]
        )
        for c, emb in zip(temp_chunks, embeddings)
    ]


def parse_xlsx(file_path: str, has_header_row: bool = True) -> List[ParsedChunkXLSX]:
    """
    Parses an Excel (.xlsx) file and returns a list of ParsedChunkXLSX objects.
    """
    temp_chunks: List[dict] = []
    workbook = openpyxl.load_workbook(file_path, data_only=True)

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]

        headers: List[str] = []
        rows = list(sheet.iter_rows())
        if not rows:
            continue

        header_row_index = 0
        if has_header_row:
            max_non_empty = 0
            for i, r in enumerate(rows[:50]):
                non_empty = sum(1 for c in r if c.value is not None and str(c.value).strip())
                if non_empty > max_non_empty:
                    max_non_empty = non_empty
                    header_row_index = i
            
            if max_non_empty > 0:
                header_row = rows[header_row_index]
                headers = [
                    str(cell.value).strip() if cell.value is not None and str(cell.value).strip() else get_column_letter(cell.column)
                    for cell in header_row
                ]

        for row_idx, row in enumerate(rows):
            row_number: int = row[0].row

            cell_details: List[dict[str, Any]] = []
            text_parts: List[str] = []

            for cell in row:
                if cell.value is None or str(cell.value).strip() == "":
                    continue

                col_index = cell.column - 1
                cell_ref = f"{get_column_letter(cell.column)}{row_number}"

                # Apply header names only for data rows (rows below the header row)
                if headers and row_idx > header_row_index and col_index < len(headers):
                    column_name = headers[col_index]
                else:
                    column_name = get_column_letter(cell.column)

                str_value = str(cell.value).strip()

                cell_details.append({
                    "ref": cell_ref,
                    "column": column_name,
                    "value": str_value,
                })
                text_parts.append(f"{column_name}: {str_value}")

            if not cell_details:
                continue

            content = " | ".join(text_parts)
            temp_chunks.append({
                "content": content,
                "location_metadata": {
                    "sheet": sheet_name,
                    "row": row_number,
                    "cells": cell_details,
                }
            })

    workbook.close()

    if not temp_chunks:
        return []

    # Batch process all collected texts
    texts = [c["content"] for c in temp_chunks]
    embeddings = get_embeddings(texts)

    return [
        ParsedChunkXLSX(
            content=c["content"],
            embedding=emb,
            location_metadata=c["location_metadata"]
        )
        for c, emb in zip(temp_chunks, embeddings)
    ]

def parse_image(file_path: str) -> List[ParsedTextFromImages]:
    """
    Extracts text from an image file using a Fireworks AI vision model.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Image not found: {file_path}")

    mime_type, _ = mimetypes.guess_type(file_path)
    if not mime_type or not mime_type.startswith("image/"):
        raise ValueError(f"Cannot determine image MIME type for: {file_path}")

    with Image.open(file_path) as img:
        width, height = img.size
        image_format = img.format or "UNKNOWN"

    with open(file_path, "rb") as image_file:
        image_base64 = base64.b64encode(image_file.read()).decode("utf-8")

    client = OpenAI(base_url=settings.VISION_MODEL_URL, api_key=settings.VISION_MODEL_API_KEY)
    response = client.chat.completions.create(
        model=settings.VISION_MODEL_NAME,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "You are an OCR assistant. Extract ALL text visible in this image. "
                            "Return each line of text on its own separate line, preserving the "
                            "original reading order (top to bottom, left to right). "
                            "Do NOT add any commentary, headings, or formatting - only the raw extracted text lines."
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:{mime_type};base64,{image_base64}"
                        },
                    },
                ],
            }
        ],
    )

    raw_text: str = response.choices[0].message.content or ""
    image_name = os.path.basename(file_path)
    
    temp_chunks: List[dict] = []

    for line_number, line in enumerate(raw_text.splitlines(), start=1):
        line = line.strip()
        if not line:
            continue

        temp_chunks.append({
            "content": line,
            "location_metadata": {
                "image_name": image_name,
                "image_path": file_path,
                "line": line_number,
                "image_format": image_format,
                "image_dimensions": {"width": width, "height": height},
            }
        })

    if not temp_chunks:
        return []

    # Batch process all collected texts
    texts = [c["content"] for c in temp_chunks]
    embeddings = get_embeddings(texts)

    return [
        ParsedTextFromImages(
            content=c["content"],
            embedding=emb,
            location_metadata=c["location_metadata"]
        )
        for c, emb in zip(temp_chunks, embeddings)
    ]